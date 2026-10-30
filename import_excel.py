#!/usr/bin/env python3
"""Download and import NZ Lotto historical data from official Excel file."""

import contextlib
import sqlite3
import sys
from pathlib import Path

import requests

DB_PATH = Path.cwd() / "data" / "lotto.db"
EXCEL_URL = "https://assets.mylotto.co.nz/assets/uploads/45c6bb20-5596-11ee-b196-005056819165.xlsx"
EXCEL_PATH = Path.cwd() / "winning_numbers.xlsx"


def download_excel() -> bool:
    if EXCEL_PATH.exists():
        print(f"Excel file already exists: {EXCEL_PATH}")
        return True

    print(f"Downloading from: {EXCEL_URL}")
    print("This may take a moment...")

    try:
        r = requests.get(EXCEL_URL, timeout=60)
        if r.status_code == 200:
            EXCEL_PATH.write_bytes(r.content)
            print(f"Downloaded: {EXCEL_PATH} ({len(r.content)} bytes)")
            return True
        else:
            print(f"Download failed: HTTP {r.status_code}")
            return False
    except Exception as e:
        print(f"Download error: {e}")
        return False


def import_excel() -> None:
    try:
        import openpyxl
    except ImportError:
        print("openpyxl not installed. Run: pip install openpyxl")
        return

    if not EXCEL_PATH.exists():
        print(f"Excel file not found: {EXCEL_PATH}")
        return

    print(f"\nReading Excel file: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

    sheet = wb[wb.sheetnames[0]]
    print(f"Using sheet: {wb.sheetnames[0]}")

    # Show first 5 rows to understand structure
    print("\nFirst 5 rows:")
    for i, row in enumerate(sheet.iter_rows(values_only=True), 1):
        if i > 5:
            break
        print(f"  Row {i}: {row}")

    # Detect headers
    headers = None
    data_start = 0
    for i, row in enumerate(sheet.iter_rows(values_only=True), 1):
        if row and any(
            str(cell).lower() in ["draw", "date", "number", "bonus", "powerball"]
            for cell in row
            if cell
        ):
            headers = [str(c).lower().strip() if c else "" for c in row]
            data_start = i
            print(f"\nHeaders at row {i}: {headers}")
            break

    if not headers:
        print("Could not detect headers.")
        return

    # Import
    conn = sqlite3.connect(str(DB_PATH.resolve()), timeout=10)
    existing_dates = {r[0] for r in conn.execute("SELECT draw_date FROM draws")}
    conn.close()

    inserted = skipped = 0

    for i, row in enumerate(sheet.iter_rows(values_only=True), 1):
        if i <= data_start:
            continue

        try:
            # Simple parsing: assume columns are Draw, Date, N1, N2, N3, N4, N5, N6, Bonus, PB
            if len(row) < 7:
                continue

            draw_num = int(row[0]) if row[0] else None
            date_val = row[1]

            if hasattr(date_val, "strftime"):
                date_str = date_val.strftime("%Y-%m-%d")
            else:
                # Try parsing string date
                from datetime import datetime

                for fmt in ["%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y"]:
                    try:
                        date_str = datetime.strptime(str(date_val).strip(), fmt).strftime(
                            "%Y-%m-%d"
                        )
                        break
                    except ValueError:
                        continue
                else:
                    continue

            if date_str in existing_dates:
                skipped += 1
                continue

            numbers = []
            for j in range(2, 8):
                if j < len(row) and row[j] is not None:
                    with contextlib.suppress(ValueError, TypeError):
                        numbers.append(int(row[j]))

            if len(numbers) != 6:
                continue

            bonus = int(row[8]) if len(row) > 8 and row[8] is not None else 0
            powerball = int(row[9]) if len(row) > 9 and row[9] is not None else 0

            conn = sqlite3.connect(str(DB_PATH.resolve()), timeout=10)
            try:
                conn.execute(
                    "INSERT INTO draws (draw_number, draw_date, numbers, bonus, powerball) VALUES (?, ?, ?, ?, ?)",
                    (draw_num, date_str, ",".join(map(str, sorted(numbers))), bonus, powerball),
                )
                conn.commit()
                inserted += 1
                if inserted % 100 == 0:
                    print(f"  ... {inserted} inserted")
            except sqlite3.IntegrityError:
                skipped += 1
            finally:
                conn.close()

        except Exception:
            continue

    print(f"\nDone. Inserted: {inserted} | Skipped: {skipped}")

    conn = sqlite3.connect(str(DB_PATH.resolve()), timeout=10)
    count = conn.execute("SELECT COUNT(*) FROM draws").fetchone()[0]
    min_date, max_date = conn.execute("SELECT MIN(draw_date), MAX(draw_date) FROM draws").fetchone()
    conn.close()
    print(f"Total draws in DB: {count}")
    if min_date and max_date:
        print(f"Date range: {min_date} to {max_date}")


def main() -> int:
    if download_excel():
        import_excel()
    return 0


if __name__ == "__main__":
    sys.exit(main())
